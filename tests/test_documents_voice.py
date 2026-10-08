"""Week 8: PDFs, Word files and OCR in the document index; the voice server's OpenAI API.
No models, no cluster; the real-OCR test runs only where Tesseract is installed (CI)."""

import io
import shutil
import wave

import pytest
from fastapi.testclient import TestClient
from qdrant_client import QdrantClient

import rag_index
import rag_server
import voice_server
from test_rag import WordEmbedding, search


def pdf_with_text(lines: list[str]) -> bytes:
    """A minimal one-page PDF whose text layer holds `lines` (Helvetica)."""
    stream = "BT /F1 14 Tf 72 720 Td 18 TL " + " ".join(f"({t}) Tj T*" for t in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


def test_pdf_text_layer_is_indexed_and_cited_by_page(tmp_path):
    (tmp_path / "invoice.pdf").write_bytes(pdf_with_text(["Invoice 2026-114 from Contoso", "Total due 349 AED by 30 October"]))
    client, embed = QdrantClient(":memory:"), WordEmbedding()
    assert rag_index.run(tmp_path, client, embed)["indexed"] == 1
    top = search(client, embed, "invoice total due contoso", k=1)
    assert top[0].node.metadata["page"] == 1 and "349 AED" in top[0].node.get_content()
    assert rag_server.format_results(top).startswith("[1] invoice.pdf p.1 (score")


def test_scanned_pdf_page_goes_through_ocr(tmp_path, monkeypatch):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(612, 792)  # no text layer, like a scan
    buf = io.BytesIO()
    w.write(buf)
    (tmp_path / "scan.pdf").write_bytes(buf.getvalue())
    import pdf2image

    monkeypatch.setattr(pdf2image, "convert_from_path", lambda *a, **k: ["page image"])
    monkeypatch.setattr(rag_index, "ocr", lambda image: "Lease signed for flat 12, rent 5000 AED")
    assert rag_index.extract(tmp_path / "scan.pdf") == [(1, "Lease signed for flat 12, rent 5000 AED")]


def test_word_file_paragraphs_and_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Project kickoff notes")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Owner", "Nixon"
    d.save(tmp_path / "kickoff.docx")
    [(page, text)] = rag_index.extract(tmp_path / "kickoff.docx")
    assert page is None and "Project kickoff notes" in text and "Owner | Nixon" in text


@pytest.mark.skipif(not shutil.which("tesseract"), reason="Tesseract not installed (CI installs it)")
def test_image_ocr_with_real_tesseract(tmp_path):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (900, 160), "white")
    ImageDraw.Draw(img).text((20, 40), "Parking permit 2026", fill="black", font=ImageFont.load_default(size=56))
    img.save(tmp_path / "permit.png")
    [(page, text)] = rag_index.extract(tmp_path / "permit.png")
    assert page is None and "Parking" in text and "2026" in text


def test_sources_line_cites_pages():
    from main import sources

    ctx = "[1] scans/invoice.pdf p.2 (score 0.70)\nTotal due.\n\n[2] notes/a.md:1-4 (score 0.66)\nMore."
    assert sources(ctx, "349 AED.") == "\n\nSources: scans/invoice.pdf p.2, notes/a.md:1-4"


def wav_bytes(seconds: float = 0.2) -> bytes:
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(16000)
        w.writeframes(b"\x00\x00" * int(16000 * seconds))
    return buf.getvalue()


def test_voice_api_speaks_and_transcribes(monkeypatch):
    seen = {}
    monkeypatch.setattr(voice_server, "synthesize", lambda text, voice, speed: seen.setdefault("tts", (text, voice)) and wav_bytes())
    monkeypatch.setattr(voice_server, "transcribe_file", lambda path, language: ("hello there", 0.2, "en"))
    api = TestClient(voice_server.app)

    r = api.post("/v1/audio/speech", json={"input": "Good morning", "voice": "alloy"})
    assert r.status_code == 200 and r.headers["content-type"] == "audio/wav" and r.content[:4] == b"RIFF"
    assert seen["tts"] == ("Good morning", voice_server.PIPER_VOICE)  # OpenAI voice names map to Piper's
    for name in ("../../etc/x-y", "en_US-other-medium"):  # CodeQL py/path-injection: never a file name
        seen.clear()
        api.post("/v1/audio/speech", json={"input": "Hi", "voice": name})
        assert seen["tts"] == ("Hi", voice_server.PIPER_VOICE)

    r = api.post("/v1/audio/transcriptions", files={"file": ("a.wav", wav_bytes(), "audio/wav")}, data={"model": "whisper-1"})
    assert r.json() == {"text": "hello there"}
    r = api.post("/v1/audio/transcriptions", files={"file": ("a.wav", wav_bytes(), "audio/wav")},
                 data={"response_format": "text"})
    assert r.text == "hello there"
    assert api.post("/v1/audio/speech", json={"input": "  "}).status_code == 400
