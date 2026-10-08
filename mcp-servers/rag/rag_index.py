"""Index the shared folder into Qdrant (LlamaIndex): read text files, PDFs, Word files and images,
split into chunks, embed with `embed-default` through LiteLLM, store with their file and line
range (or PDF page). PDF pages without a text layer (scans) and images are read with Tesseract
OCR (Week 8, README §6.18).

Incremental: every chunk carries its file's sha256, so a run re-embeds only new or changed files
and drops the chunks of deleted ones. Runs as the k3s CronJob `rag-index` (`make rag-index` for
a run now); Dagster takes it over in Week 4.

    python rag_index.py            # index FILES_ROOT
    python rag_index.py --rebuild  # drop the collection first (after changing model or chunking)
"""

import argparse
import hashlib
import logging
import os
import time
from pathlib import Path

from llama_index.core import Document
from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.schema import MetadataMode
from qdrant_client import QdrantClient

from rag_store import COLLECTION, QDRANT_URL, LiteLLMEmbedding, vector_store

log = logging.getLogger("rag-index")

ROOT = Path(os.environ.get("FILES_ROOT", "/data")).resolve()
# Small chunks: a 3B model answers faster and better from 3–4 short passages than from one long one
CHUNK_TOKENS = int(os.environ.get("CHUNK_TOKENS", "256"))
CHUNK_OVERLAP = int(os.environ.get("CHUNK_OVERLAP", "32"))
MAX_FILE_BYTES = 1_000_000  # text files
MAX_DOC_BYTES = 25_000_000  # PDFs, Word files, images
TEXT_SUFFIXES = {
    ".md", ".txt", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml", ".ini", ".log",
    ".py", ".sh", ".ps1", ".sql", ".html", ".xml",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
DOC_SUFFIXES = {".pdf", ".docx"} | IMAGE_SUFFIXES
OCR_LANG = os.environ.get("OCR_LANG", "eng")  # Tesseract languages, e.g. "eng+ara"
OCR_MIN_CHARS = 30  # a PDF page with less text than this is treated as a scan


def scan(root: Path) -> dict[str, tuple[Path, str]]:
    """Indexable files under root: relative path -> (path, sha256). Skips big, unreadable,
    unsupported and outside files."""
    files = {}
    for f in sorted(root.rglob("*")):
        suffix = f.suffix.lower()
        if not f.is_file() or suffix not in TEXT_SUFFIXES | DOC_SUFFIXES or not f.resolve().is_relative_to(root):
            continue
        if f.stat().st_size > (MAX_FILE_BYTES if suffix in TEXT_SUFFIXES else MAX_DOC_BYTES):
            continue
        data = f.read_bytes()
        if suffix in TEXT_SUFFIXES:
            try:
                data.decode("utf-8")
            except UnicodeDecodeError:
                continue
        files[f.relative_to(root).as_posix()] = (f, hashlib.sha256(data).hexdigest())
    return files


def ocr(image) -> str:
    """Tesseract OCR of a PIL image; "" (with a warning) when Tesseract isn't installed."""
    import pytesseract

    try:
        return pytesseract.image_to_string(image, lang=OCR_LANG)
    except pytesseract.TesseractNotFoundError:
        log.warning("Tesseract not installed: skipping OCR")
        return ""


def extract(f: Path) -> list[tuple[int | None, str]]:
    """(page, text) pieces of one file: one per PDF page, one (page None) for anything else."""
    suffix = f.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return [(None, f.read_text(encoding="utf-8"))]
    if suffix == ".docx":
        import docx

        d = docx.Document(str(f))
        parts = [p.text for p in d.paragraphs]
        parts += [" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows]
        return [(None, "\n".join(parts))]
    if suffix in IMAGE_SUFFIXES:
        from PIL import Image

        with Image.open(f) as img:
            return [(None, ocr(img))]
    if suffix == ".pdf":
        from pypdf import PdfReader

        pages = []
        for i, page in enumerate(PdfReader(str(f)).pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < OCR_MIN_CHARS:  # no text layer: a scan, read it with OCR
                from pdf2image import convert_from_path

                images = convert_from_path(str(f), dpi=200, first_page=i, last_page=i)
                text = ocr(images[0]) if images else text
            pages.append((i, text))
        return pages
    return []


def indexed(client: QdrantClient) -> dict[str, str]:
    """What the collection holds now: relative path -> sha256."""
    if not client.collection_exists(COLLECTION):
        return {}
    seen, offset = {}, None
    while True:
        points, offset = client.scroll(
            COLLECTION, limit=1000, offset=offset, with_payload=["file_path", "file_hash"], with_vectors=False
        )
        for p in points:
            seen[p.payload["file_path"]] = p.payload["file_hash"]
        if offset is None:
            return seen


def plan(current: dict[str, str], stored: dict[str, str]) -> tuple[list[str], list[str]]:
    """(paths to (re)index, paths to delete). Changed files are in both lists."""
    to_index = sorted(p for p, h in current.items() if stored.get(p) != h)
    to_delete = sorted(p for p in stored if p not in current or p in to_index)
    return to_index, to_delete


def chunks(path: str, text: str, file_hash: str, splitter: SentenceSplitter, page: int | None = None):
    """Split one file (or one PDF page) into nodes that know their line range and page. The doc
    id is the path, so a file's chunks can be deleted together."""
    metadata = {"file_path": path, "file_hash": file_hash}
    if page is not None:
        metadata["page"] = page
    doc = Document(text=text, id_=path, metadata=metadata)
    # The path goes into the embedded text (it helps retrieval); hash and lines don't
    doc.excluded_embed_metadata_keys = ["file_hash", "lines"]
    doc.excluded_llm_metadata_keys = ["file_hash"]
    nodes = splitter.get_nodes_from_documents([doc])
    for n in nodes:
        start = n.start_char_idx if n.start_char_idx is not None else max(text.find(n.text), 0)
        end = n.end_char_idx if n.end_char_idx is not None else start + len(n.text)
        n.metadata["lines"] = f"{text.count(chr(10), 0, start) + 1}-{text.count(chr(10), 0, max(end - 1, start)) + 1}"
    return nodes


def run(root: Path = ROOT, client: QdrantClient | None = None, embed_model=None, rebuild: bool = False) -> dict:
    client = client or QdrantClient(url=QDRANT_URL)
    embed_model = embed_model or LiteLLMEmbedding()
    store = vector_store(client)
    if rebuild and client.collection_exists(COLLECTION):
        client.delete_collection(COLLECTION)

    current = scan(root)
    to_index, to_delete = plan({p: h for p, (_, h) in current.items()}, indexed(client))
    for p in to_delete:
        store.delete(ref_doc_id=p)

    splitter = SentenceSplitter(chunk_size=CHUNK_TOKENS, chunk_overlap=CHUNK_OVERLAP)
    n_chunks = 0
    for p in to_index:
        f, file_hash = current[p]
        try:
            pieces = extract(f)
        except Exception as e:  # a broken PDF must not stop the whole run
            log.warning("skipping %s: %s", p, e)
            continue
        nodes = [n for page, text in pieces if text.strip() for n in chunks(p, text, file_hash, splitter, page)]
        if not nodes:
            continue
        vectors = embed_model.get_text_embedding_batch([n.get_content(MetadataMode.EMBED) for n in nodes])
        for n, v in zip(nodes, vectors):
            n.embedding = v
        store.add(nodes)
        n_chunks += len(nodes)
        log.info("indexed %s: %d chunks", p, len(nodes))

    removed = [p for p in to_delete if p not in to_index]
    return {"files": len(current), "indexed": len(to_index), "chunks": n_chunks, "removed": len(removed),
            "unchanged": len(current) - len(to_index)}


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)  # one line per Qdrant/LiteLLM call is noise
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--rebuild", action="store_true", help="drop the collection and index everything")
    args = ap.parse_args()
    t = time.time()
    stats = run(rebuild=args.rebuild)
    log.info("done in %.1fs: %s (collection %s)", time.time() - t, stats, COLLECTION)
